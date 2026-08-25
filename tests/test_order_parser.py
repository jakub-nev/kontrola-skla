# tests/test_order_parser.py
from pathlib import Path

import pytest

from order_parser import parse_order, parse_skladba

SAMPLE = Path(__file__).resolve().parent.parent / "Glass_solutions" / "objednavka_sklo_VltavaHolding_16.9.25.xlsx"


# --- parse_skladba ---------------------------------------------------------

def test_skladba_triple():
    assert parse_skladba("4-18-4-18-4") == ([4.0, 4.0, 4.0], [18.0, 18.0])


def test_skladba_double():
    assert parse_skladba("4-16-4") == ([4.0, 4.0], [16.0])


def test_skladba_laminated_with_suffixes():
    panes, gaps = parse_skladba("33.2 XN-16-4-16-6 XN")
    assert panes == [33.2, 4.0, 6.0]
    assert gaps == [16.0, 16.0]


def test_skladba_unparseable_returns_none():
    assert parse_skladba("trojsklo standard") == (None, None)
    assert parse_skladba("") == (None, None)
    assert parse_skladba("4-18") == (None, None)  # even token count = ambiguous


# --- parse_order on the real sample file -----------------------------------

@pytest.fixture(scope="module")
def items():
    return parse_order(SAMPLE)


def test_row_count(items):
    assert len(items) == 11  # trailing empty row ignored


def test_first_row(items):
    it = items[0]
    assert (it.objekt, it.width, it.height, it.quantity) == ("becica", 830, 1400, 1)
    assert it.skladba_raw == "4-18-4-18-4"
    assert it.panes == [4.0, 4.0, 4.0]
    assert it.gaps == [18.0, 18.0]


def test_laminated_row(items):
    it = items[7]  # Položka 8, RD_zelechovice
    assert (it.width, it.height, it.quantity) == (654, 1140, 6)
    assert it.panes == [33.2, 4.0, 6.0]
    assert it.gaps == [16.0, 16.0]
    assert "protihlukov" in it.typ.lower()


def test_missing_column_raises(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Položka", "Objekt"])  # no width/height/kus
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    with pytest.raises(ValueError, match="chybí sloupec"):
        parse_order(bad)


def test_zero_quantity_preserved(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Položka", "Objekt", "Označení pozice", "Šířka (mm)",
                      "Výška (mm)", "Kus", "Skladba skla", "Typ skla"])
    wb.active.append([1, "x", 1, 800, 600, 0, "4-16-4", "dvojsklo"])
    p = tmp_path / "zero.xlsx"
    wb.save(p)
    assert parse_order(p)[0].quantity == 0


def test_kusovnik_column_not_mistaken_for_kus(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Kusovník", "Šířka (mm)", "Výška (mm)", "Kus"])
    wb.active.append(["A-1", 800, 600, 3])
    p = tmp_path / "kusovnik.xlsx"
    wb.save(p)
    assert parse_order(p)[0].quantity == 3
